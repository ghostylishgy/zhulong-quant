from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import asdict, dataclass
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parents[1]
SRC_DIR = BASE_DIR / "src"
if str(SRC_DIR) not in sys.path:
    sys.path.insert(0, str(SRC_DIR))

from spark_etf.services.nav_service import NavService
from spark_etf.services.weekly_snapshot_service import COST_BASIS_TYPES, WeeklySnapshotService

MARKET_VALUE_RE = re.compile(r"(?:当前市值|市值|market_value)\s*[:：=]?\s*([-+]?\d+(?:,\d{3})*(?:\.\d+)?)", re.IGNORECASE)


@dataclass(frozen=True, slots=True)
class MarketValueOverride:
    fund_code: str
    fund_name: str
    shares: float
    market_value: float
    source: str


def _parse_snapshot_date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y%m%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"Invalid snapshot date: {value!r}")


def _as_float(value: object) -> float | None:
    if value is None:
        return None
    text = str(value).strip().replace(",", "")
    if not text:
        return None
    return float(text)


def _normalize_code(value: object) -> str:
    code = str(value or "").strip()
    if not code:
        return ""
    if "." not in code and len(code) == 6 and code.isdigit():
        return f"{code}.OF"
    return code


def _market_value_from_note(note: str) -> float | None:
    match = MARKET_VALUE_RE.search(note)
    if not match:
        return None
    return float(match.group(1).replace(",", ""))


def workbook_to_text(path: Path) -> tuple[date, str, str, list[MarketValueOverride]]:
    wb = load_workbook(path, data_only=False)
    if "本周持仓" not in wb.sheetnames:
        raise ValueError("Workbook must contain sheet '本周持仓'")
    ws = wb["本周持仓"]
    snapshot_date = _parse_snapshot_date(ws["B2"].value)
    default_cost_basis = str(ws["B3"].value or "tiantian_position_cost").strip()
    if default_cost_basis not in COST_BASIS_TYPES:
        raise ValueError(f"Invalid workbook cost basis type: {default_cost_basis}")

    lines: list[str] = []
    market_value_overrides: list[MarketValueOverride] = []
    for row_idx in range(8, 38):
        code = _normalize_code(ws.cell(row_idx, 1).value)
        name = str(ws.cell(row_idx, 2).value or "").strip()
        shares = _as_float(ws.cell(row_idx, 3).value)
        unit_cost = _as_float(ws.cell(row_idx, 4).value)
        cost_basis = str(ws.cell(row_idx, 5).value or default_cost_basis).strip()
        note = str(ws.cell(row_idx, 7).value or "").strip()
        confirmed = str(ws.cell(row_idx, 8).value or "是").strip()
        market_value = _as_float(ws.cell(row_idx, 9).value)

        if not code and not name and shares is None and unit_cost is None:
            continue
        if confirmed != "是":
            continue
        if not code:
            raise ValueError(f"Row {row_idx}: missing fund code")
        if not name:
            name = code
        if shares is None:
            raise ValueError(f"Row {row_idx}: missing shares")
        if unit_cost is None:
            raise ValueError(f"Row {row_idx}: missing unit cost")
        if cost_basis != default_cost_basis:
            raise ValueError(f"Row {row_idx}: mixed cost basis is not supported in one import batch")

        if market_value is None and note:
            market_value = _market_value_from_note(note)
        if market_value is not None:
            if shares <= 0:
                raise ValueError(f"Row {row_idx}: market value override requires positive shares")
            market_value_overrides.append(
                MarketValueOverride(
                    fund_code=code,
                    fund_name=name,
                    shares=shares,
                    market_value=market_value,
                    source="xlsx_column_i" if ws.cell(row_idx, 9).value is not None else "note_marker",
                )
            )

        note_text = f" 备注 {note}" if note else ""
        lines.append(f"{code} {name} 份额 {shares:.8f} 成本 {unit_cost:.8f}{note_text}")

    return snapshot_date, default_cost_basis, "\n".join(lines), market_value_overrides


def upsert_market_value_overrides(
    snapshot_date: date,
    overrides: list[MarketValueOverride],
    db_path: str | None,
) -> list[dict[str, object]]:
    if not overrides:
        return []

    nav_service = NavService(db_path=db_path)
    payloads = []
    result = []
    for item in overrides:
        unit_nav = item.market_value / item.shares
        payloads.append(
            {
                "fund_code": item.fund_code,
                "nav_date": snapshot_date,
                "unit_nav": unit_nav,
                "adj_nav": unit_nav,
                "source": "manual_market_value",
                "data_quality": "DATA_OK",
                "raw": {
                    "source": item.source,
                    "fund_name": item.fund_name,
                    "market_value": item.market_value,
                    "shares": item.shares,
                    "calculation": "unit_nav = market_value / shares",
                },
            }
        )
        result.append(
            {
                **asdict(item),
                "nav_date": snapshot_date.isoformat(),
                "unit_nav": unit_nav,
            }
        )
    nav_service.upsert_nav_rows(payloads)
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description="Import Spark ETF weekly snapshot from the XLSX template")
    parser.add_argument("--xlsx", required=True, type=Path)
    parser.add_argument("--db", type=str, default=None)
    args = parser.parse_args()

    snapshot_date, cost_basis_type, text, market_value_overrides = workbook_to_text(args.xlsx)
    service = WeeklySnapshotService(db_path=args.db)
    result = service.import_text(
        snapshot_date=snapshot_date,
        text=text,
        cost_basis_type=cost_basis_type,
        source="manual_xlsx",
        source_path=str(args.xlsx),
        note="imported from weekly xlsx template",
    )
    result["market_value_overrides"] = upsert_market_value_overrides(snapshot_date, market_value_overrides, args.db)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
